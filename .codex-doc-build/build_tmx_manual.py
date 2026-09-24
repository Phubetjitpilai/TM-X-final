from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path(r'D:\All Work\TM-X React\output\TM-X_Web_User_Manual_Template.docx')
out.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = Inches(.75)
sec.bottom_margin = Inches(.68)
sec.left_margin = Inches(.85)
sec.right_margin = Inches(.85)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Tahoma'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0,0,0)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.18
for name,size,before,after in [('Title',22,0,14),('Heading 1',15,20,9),('Heading 2',12,12,6)]:
    st=styles[name]
    st.font.name='Tahoma'; st.font.size=Pt(size); st.font.bold=True
    st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(before)
    st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True

if 'Placeholder' not in styles:
    st=styles.add_style('Placeholder',WD_STYLE_TYPE.PARAGRAPH)
else: st=styles['Placeholder']
st.font.name='Tahoma'; st.font.size=Pt(10); st.font.italic=True
st.font.color.rgb=RGBColor(95,108,121)
st.paragraph_format.space_before=Pt(4)
st.paragraph_format.space_after=Pt(10)

def p(text='',style=None):
    return doc.add_paragraph(text,style)
def h(text):
    doc.add_heading(text,level=1)
def numbered(text):
    z=p(style='List Number'); z.add_run(text)
def bullet(text):
    z=p(style='List Bullet'); z.add_run(text)
def image(label):
    p(f'ภาพประกอบ  [วางภาพหน้าจอ {label}]','Placeholder')
def field(paragraph, name):
    run=paragraph.add_run()
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    inst=OxmlElement('w:instrText'); inst.set(qn('xml:space'),'preserve'); inst.text=name
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate')
    txt=OxmlElement('w:t'); txt.text='1'
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    for el in [begin,inst,sep,txt,end]: run._r.append(el)

header=sec.header.paragraphs[0]
header.text='TM-X Control System'
header.style=styles['Normal']
header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('หน้า ').font.size=Pt(8)
field(footer,'PAGE')

doc.add_paragraph('คู่มือการใช้งานเว็บ TM-X Control System','Title')
p('สำหรับผู้ปฏิบัติงานและผู้ดูแลข้อมูล')
p('หน่วยงาน  [ระบุ]    เวอร์ชัน  [ระบุ]    วันที่  [ระบุ]','Placeholder')
p('เอกสารนี้จัดลำดับการใช้งานตั้งแต่การตรวจสอบความพร้อม การเตรียมชิ้นงาน การวัดและดูผล ไปจนถึงการแก้ไขและส่งออกข้อมูล เติมรายละเอียดและภาพหน้าจอในตำแหน่งที่ระบุเพื่อให้ตรงกับขั้นตอนปฏิบัติงานจริง')

h('1 ภาพรวมการใช้งาน')
p('เว็บ TM-X Control System มีหน้าหลักสำหรับควบคุมการวัด ดูผล แก้ไขข้อมูล และส่งออกผลการวัด')
bullet('Home  เตรียมข้อมูลชิ้นงาน เริ่มหรือหยุดการวัด และติดตามผล')
bullet('Edit  ค้นหา เพิ่ม หรือแก้ไขข้อมูลที่เกี่ยวข้อง')
bullet('Export  เลือกรูปแบบรายงานและดาวน์โหลดข้อมูล')
image('เมนูหลักของเว็บ')

h('2 ก่อนเริ่มใช้งาน')
numbered('ตรวจสอบว่าเชื่อมต่ออะแดปเตอร์ Ethernet กับแล็ปท็อปเรียบร้อยแล้ว')
numbered('เปิดเว็บเบราว์เซอร์และไปที่ http://192.168.10.10:8000')
numbered('ตรวจสอบสถานะ Session: Stopped ซึ่งหมายถึงระบบหยุดการวัดและรอเริ่มงาน')
numbered('ตรวจสอบสถานะ Raspberry Pi: Online ซึ่งหมายถึง Raspberry Pi เชื่อมต่อกับ PC Server')
numbered('ตรวจสอบสถานะ Server: Online ซึ่งหมายถึงหน้าเว็บเชื่อมต่อกับ PC Server และรับส่งข้อมูลได้')
image('หน้า Home และป้ายสถานะ')
p('[เพิ่มวิธีปฏิบัติเมื่อสถานะไม่เป็นไปตามที่ระบุ]','Placeholder')

h('3 เตรียมข้อมูลชิ้นงาน')
numbered('[เลือกประเภทการวัด IPM, New หรือ Rework]')
numbered('[กรอกชื่อ Operator และข้อมูล ALPL]')
numbered('[ตรวจสอบรายการในคิวแล้วกด Save]')
image('ส่วน Part Entry')
p('[เพิ่มตัวอย่างข้อมูลที่กรอกถูกต้องและเงื่อนไขก่อนเริ่มวัด]','Placeholder')

h('4 เริ่มและหยุดการวัด')
numbered('[ตรวจสอบข้อมูลชิ้นงานและสถานะระบบก่อนกด Start]')
numbered('[กด Start และตรวจสอบว่า Session เปลี่ยนเป็น Running]')
numbered('[กด Stop เมื่อต้องการสิ้นสุดการวัด]')
image('ส่วน Session Control')
p('[อธิบายผลของการหยุดการวัดกลางคัน]','Placeholder')

h('5 อ่านผลระหว่างวัด')
numbered('[อ่านค่า X, Y และ Offset ตามประเภทการวัด]')
numbered('[ตรวจสอบผล OK หรือ NG และภาพล่าสุดจาก Camera Preview]')
numbered('[ติดตามจำนวนชิ้นงานที่วัดแล้วใน Progress]')
image('ส่วน Live Telemetry และ Camera Preview')
p('[เพิ่มเกณฑ์การตัดสินผลและตัวอย่างที่พบจริง]','Placeholder')

h('6 ค้นหาและดูผลย้อนหลัง')
numbered('[ค้นหารายการด้วย ALPL หรือเลือกวันที่]')
numbered('[เปิดรายการ Measurements ที่ต้องการดู]')
numbered('[ตรวจสอบค่าที่วัด ผลการตัดสิน และภาพประกอบ]')
image('ตาราง Measurements และหน้ารายละเอียด')

h('7 แก้ไขข้อมูล')
numbered('[เลือกตาราง Parts หรือ Measurements ในหน้า Edit]')
numbered('[เพิ่มหรือแก้ไขข้อมูลที่ต้องการ แล้วกด Save]')
numbered('[ตรวจสอบข้อมูลหลังบันทึก]')
image('หน้า Edit')
p('[ระบุผู้มีสิทธิ์แก้ไขข้อมูลและข้อควรระวัง]','Placeholder')

h('8 ส่งออกข้อมูล')
numbered('[เลือกรูปแบบ CSV, PDF หรือ Excel]')
numbered('[เลือก Template และกำหนดตัวกรองข้อมูล]')
numbered('[ตรวจตัวอย่างข้อมูล แล้วกดดาวน์โหลด]')
image('หน้า Export')
p('[ระบุรูปแบบการตั้งชื่อไฟล์และตำแหน่งจัดเก็บ]','Placeholder')

h('9 ปัญหาที่พบบ่อย')
for name,desc in [
    ('Raspberry Pi ไม่เชื่อมต่อ','[วิธีตรวจสอบและผู้รับผิดชอบ]'),
    ('กด Start ไม่ได้','[ตรวจสอบข้อมูลคิวและสถานะระบบ]'),
    ('ไม่มีภาพหรือผลวัด','[วิธีตรวจสอบเครื่องวัดและการรับข้อมูล]'),
    ('Export แล้วไม่มีข้อมูล','[ตรวจสอบตัวกรองและช่วงวันที่]'),
]:
    doc.add_heading(name,level=2); p(desc,'Placeholder')

h('10 คำศัพท์และผู้ติดต่อ')
for name in ['ALPL','Session','Tolerance','OK / NG']:
    z=p(); z.add_run(name+'  ').bold=True; z.add_run('[ความหมาย]')
doc.add_heading('ผู้ดูแลระบบ',level=2)
p('ชื่อ  [ระบุ]\nแผนก  [ระบุ]\nโทรศัพท์หรืออีเมล  [ระบุ]')
p('เอกสารอ้างอิง  [ระบุ]','Placeholder')

doc.save(out)
print(out)
