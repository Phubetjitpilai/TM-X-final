from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(r'D:\All Work\TM-X React')
IMG=ROOT/'.codex-doc-build'/'source_images'
OUT=ROOT/'output'/'TM-X_Web_User_Manual_Revised.docx'
OUT.parent.mkdir(exist_ok=True)

d=Document()
s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=Inches(.72); s.bottom_margin=Inches(.66)
s.left_margin=Inches(.82); s.right_margin=Inches(.82)
s.header_distance=Inches(.28); s.footer_distance=Inches(.3)

styles=d.styles
normal=styles['Normal']; normal.font.name='Tahoma'; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1.16; normal.paragraph_format.space_after=Pt(5)
for name,size,before,after in [('Title',22,0,12),('Heading 1',16,16,8),('Heading 2',12,11,5)]:
    st=styles[name]; st.font.name='Tahoma'; st.font.size=Pt(size); st.font.bold=True
    st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True
for name,size,color,italic in [('Caption',8.5,RGBColor(84,98,110),False),('Small Note',9.5,RGBColor(60,68,77),False)]:
    st=styles[name] if name in styles else styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st.font.name='Tahoma'; st.font.size=Pt(size); st.font.color.rgb=color; st.font.italic=italic
    st.paragraph_format.space_after=Pt(6)

def para(text='',style=None): return d.add_paragraph(text,style)
def h1(text): d.add_heading(text,level=1)
def h2(text): d.add_heading(text,level=2)
def step(n,text):
    q=d.add_paragraph()
    q.paragraph_format.left_indent=Inches(.23)
    q.paragraph_format.first_line_indent=Inches(-.23)
    q.paragraph_format.space_after=Pt(6)
    q.add_run(f'{n}.  ').bold=True
    q.add_run(text)
    return q
def bullet(text):
    q=d.add_paragraph(style='List Bullet'); q.add_run(text); return q
def note(label,text):
    q=d.add_paragraph(style='Small Note'); q.add_run(label+'  ').bold=True; q.add_run(text); return q
def fig(key,caption,maxh=3.25,maxw=6.85):
    fn=IMG/f'{key}.png'
    if not fn.exists(): raise FileNotFoundError(fn)
    im=Image.open(fn)
    w,h=im.size
    scale=min(maxw/w,maxh/h)
    wp,hp=w*scale,h*scale
    q=d.add_paragraph()
    q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_before=Pt(3)
    q.paragraph_format.space_after=Pt(2)
    q.paragraph_format.keep_with_next=True
    q.add_run().add_picture(str(fn),width=Inches(wp),height=Inches(hp))
    import re
    caption=re.sub(r'^ภาพ\s+\d+\s+', 'ภาพประกอบ  ', caption)
    c=d.add_paragraph(caption,'Caption'); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return q
def page(): d.add_page_break()
def field(p,name):
    r=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin')
    b=OxmlElement('w:instrText'); b.set(qn('xml:space'),'preserve'); b.text=name
    c=OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'),'separate')
    t=OxmlElement('w:t'); t.text='1'
    e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),'end')
    for x in (a,b,c,t,e): r._r.append(x)

hdr=s.header.paragraphs[0]; hdr.text='TM-X Control System  คู่มือการใช้งานเว็บ'; hdr.runs[0].font.name='Tahoma'; hdr.runs[0].font.size=Pt(8)
ftr=s.footer.paragraphs[0]; ftr.alignment=WD_ALIGN_PARAGRAPH.RIGHT
ftr.add_run('หน้า ').font.size=Pt(8); field(ftr,'PAGE')

d.add_paragraph('คู่มือการใช้งานเว็บ TM-X Control System','Title')
para('สำหรับผู้ปฏิบัติงานและผู้ดูแลข้อมูล')
para('หน่วยงาน  MNT     เวอร์ชัน  [ระบุ]     วันที่ปรับปรุง  [ระบุ]','Small Note')
para('คู่มือนี้อธิบายการเชื่อมต่อระบบ การเตรียมข้อมูลชิ้นงาน การวัด การดูผล และการจัดการข้อมูลผ่านเว็บ TM-X Control System ผู้ปฏิบัติงานควรตรวจสถานะระบบก่อนเริ่มวัดทุกครั้ง และตรวจหมายเลข ALPL ในคิวให้ตรงกับชิ้นงานที่วางบนเครื่อง')

h1('1 ภาพรวมระบบ')
para('เว็บแบ่งงานออกเป็นสามหน้า ได้แก่ Measure สำหรับการวัดและดูผล Edit สำหรับจัดการข้อมูล และ Export สำหรับดาวน์โหลดข้อมูล')
h2('สถานะการเชื่อมต่อ')
bullet('Server Online หมายถึงหน้าเว็บติดต่อ PC Server ได้ และระบบพร้อมรับส่งข้อมูล')
bullet('Server Offline หมายถึงหน้าเว็บขาดการเชื่อมต่อกับ PC Server ข้อมูลสถานะที่ยังปรากฏบนจออาจเป็นค่าก่อนการเชื่อมต่อหลุด')
bullet('DB Offline หมายถึงหน้าเว็บยังติดต่อ PC Server ได้ แต่เซิร์ฟเวอร์ติดต่อฐานข้อมูลไม่ได้ จึงอาจโหลดหรือบันทึกข้อมูลไม่สำเร็จ')
fig('p007_01','ภาพ 1  หน้าหลัก Measure และตำแหน่งป้ายสถานะ',3.05)
fig('p009_01','ภาพ 2  ตัวอย่างสถานะเมื่อเซิร์ฟเวอร์ขาดการเชื่อมต่อ',.58)
fig('p011_01','ภาพ 3  ตัวอย่างสถานะเมื่อฐานข้อมูลขาดการเชื่อมต่อ',.58)
h2('หน้าหลักของเว็บ')
bullet('Measure แสดง Session Control, Part Entry, Live Telemetry, Camera Preview และ Measurement History')
bullet('Edit ใช้เพิ่ม แก้ไข หรือลบข้อมูล ALPL และผลการวัด รวมถึงจัดการตารางอ้างอิงและรายการในถังขยะ')
bullet('Export ใช้เลือกข้อมูลและส่งออกเป็น CSV, PDF หรือ Excel')

h1('2 ก่อนเริ่มใช้งาน')
step(1,'ต่ออะแดปเตอร์ Ethernet เข้ากับแล็ปท็อป และตรวจว่าสายเครือข่ายต่อแน่น')
fig('p035_01','ภาพ 11  ตัวอย่างการต่ออะแดปเตอร์ Ethernet',2.75,3.6)
step(2,'เปิดเว็บเบราว์เซอร์ แล้วไปที่ http://192.168.10.10:8000')
fig('p037_01','ภาพ 12  ที่อยู่เว็บสำหรับเข้าใช้งาน',.7,4.2)
step(3,'ตรวจสถานะ Session ให้เป็น STOPPED, Raspberry Pi ให้เป็น Online และ Server ให้เป็น Online ก่อนเตรียมการวัด')
fig('p039_01','ภาพ 13  สถานะที่ควรเห็นก่อนเริ่มงาน',1.1)
note('หากสถานะไม่พร้อม','อย่าเริ่มวัด ให้ตรวจสายและการเชื่อมต่อก่อน หากยังไม่กลับมา Online ให้แจ้งผู้ดูแลระบบ')

h1('3 เตรียมข้อมูลชิ้นงาน')
step(1,'ในหน้า Measure กด New Entry เพื่อเปิดแบบฟอร์ม Part Entry')
fig('p041_01','ภาพ 14  ปุ่ม New Entry',1.1)
step(2,'เลือกประเภทงาน IPM, New หรือ Rework ให้ตรงกับชิ้นงาน')
fig('p043_01','ภาพ 15  ตัวเลือกประเภทงานใน Part Entry',2.9)
step(3,'เลือก Operator หรือชื่อผู้ปฏิบัติงาน')
fig('p046_01','ภาพ 16  ช่อง Operator',2.35)
step(4,'เลือกวิธีสั่งวัดในช่อง Trigger: Auto (MCU) ให้ระบบสั่งวัดเมื่อ MCU ตรวจพบชิ้นงาน หรือ Manual (ปุ่มบนเว็บ) เพื่อกดปุ่ม Trigger ด้วยตนเองทุกชิ้น')
fig('p049_01','ภาพ 17  ตัวเลือก Auto และ Manual',2.35)
note('ข้อสำคัญ','เลือกวิธีสั่งวัดก่อนกด Start และตรวจว่าอุปกรณ์ที่ใช้กับโหมด Auto เชื่อมต่อพร้อมทำงาน')
step(5,'กรอกหมายเลข ALPL ตามตัวเลขบนชิ้นงาน สามารถกรอกหลายหมายเลขในกลุ่มเดียวได้')
bullet('400-403 หมายถึง ALPL 400, 401, 402 และ 403')
bullet('400,403 หมายถึง ALPL 400 และ 403')
bullet('399,400-403 หมายถึง ALPL 399 ถึง 403')
fig('p052_01','ภาพ 18  ช่องกรอกหมายเลข ALPL',1.1)
step(6,'กรอกข้อมูลอื่นให้ครบตามประเภทงานที่เลือก โดย ALPL ในกลุ่มเดียวกันต้องใช้ข้อมูลชุดเดียวกัน')
bullet('IPM ระบุ Package Size และ Handler')
bullet('New และ Rework ระบุ Package Size, Part Number, PO Number, Vendor, Owner, Description และ Receive Date')
fig('p056_01','ภาพ 19  ตัวอย่างช่องข้อมูลของงาน IPM',1.1)
fig('p058_01','ภาพ 20  ตัวอย่างช่องข้อมูลของงาน New หรือ Rework',3.0)
step(7,'หาก ALPL อีกชุดใช้ข้อมูลต่างจากกลุ่มแรก ให้กด Add Group แล้วกรอกข้อมูลของกลุ่มใหม่')
fig('p061_01','ภาพ 21  ปุ่ม Add Group',.58)
step(8,'ตรวจข้อมูลทุกกลุ่ม แล้วกด Save เพื่อบันทึกคิวการวัด')
fig('p063_01','ภาพ 22  ปุ่ม Save',.58)

h1('4 เริ่มการวัด')
fig('p015_01','ภาพ 4  หน้า Measure ส่วนควบคุมและผลล่าสุด',2.55)
step(1,'ตรวจหมายเลข ALPL และข้อมูลชิ้นงานใน Part Entry อีกครั้ง หากพบข้อผิดพลาดให้กด Edit เพื่อแก้ก่อนเริ่ม')
fig('p065_01','ภาพ 23  รายละเอียดชิ้นงานก่อนเริ่มวัด',2.65,4.8)
step(2,'กด Start แล้วตรวจว่า Session เปลี่ยนเป็น RUNNING')
fig('p069_01','ภาพ 24  สถานะ RUNNING และคิวชิ้นงาน',2.7)
step(3,'ดูกรอบสีน้ำเงินใน Queue เพื่อยืนยัน ALPL ที่ต้องวัด วางชิ้นงานหมายเลขนั้นบนตำแหน่งวัด')
fig('p074_01','ภาพ 25  Queue แสดงชิ้นงานที่กำลังรอวัด',2.65)
step(4,'ถ้าเลือก Manual ให้กด Trigger บนเว็บหนึ่งครั้งต่อชิ้น และรอผลวัดปรากฏ ถ้าเลือก Auto ระบบจะสั่งวัดเมื่อ MCU ตรวจพบชิ้นงาน')
fig('p074_02','ภาพ 26  ปุ่ม Trigger สำหรับโหมด Manual',2.7)
fig('p077_01','ภาพ 27  ตัวอย่างข้อความตอบรับหลังสั่ง Trigger',2.45)
step(5,'ตรวจค่าและภาพล่าสุดบนหน้าเว็บ จากนั้นนำชิ้นงานที่วัดแล้วออก')
fig('p081_01','ภาพ 28  ตัวอย่างผลวัด OK',2.7)
step(6,'ตรวจ Queue และ Progress แล้ววางชิ้นงานหมายเลขถัดไป ทำซ้ำจนวัดครบทุกชิ้น')
fig('p088_01','ภาพ 29  ผลใน Queue และความคืบหน้าการวัด',2.7)

h1('5 อ่านผลการวัด')
para('Live Telemetry แสดงค่าที่วัดล่าสุดและผลรวมของชิ้นงาน ส่วน Camera Preview แสดงภาพจากการวัดล่าสุด ผล OK หมายถึงค่าที่ใช้ตัดสินทุกค่าผ่านเกณฑ์ที่ตั้งไว้ และ NG หมายถึงมีอย่างน้อยหนึ่งค่าไม่ผ่านเกณฑ์')
h2('งาน IPM')
bullet('Value X คือขนาดรูเปิดในแนวนอน และ Value Y คือขนาดรูเปิดในแนวตั้ง หน่วยเป็นมิลลิเมตร')
bullet('ระบบตัดสินผลจาก Value X และ Value Y โดยทั้งสองค่าต้องอยู่ในช่วงที่ยอมรับได้จึงเป็น OK')
h2('งาน New และ Rework')
bullet('ระบบแสดง Value X, Value Y และค่า Offset Opening สำหรับการคลาดเคลื่อนของรูเปิดจากตำแหน่งอ้างอิง')
bullet('OP-X แสดงการคลาดเคลื่อนในแนวนอน และ OP-Y แสดงการคลาดเคลื่อนในแนวตั้ง')
bullet('Offset Opening Position แสดงทิศทางตำแหน่งของรูเปิดเมื่อเทียบกับจุดอ้างอิง')
bullet('ผล OK ต้องผ่านเกณฑ์ที่ตั้งไว้สำหรับ Value X, Value Y, OP-X และ OP-Y')
note('การอ่าน Queue','สีเขียวหมายถึงวัดแล้วและผ่าน สีแดงหมายถึงวัดแล้วและไม่ผ่าน กรอบสีน้ำเงินระบุชิ้นงานที่กำลังเลือก ส่วน Progress แสดงจำนวนที่วัดแล้วและจำนวน OK/NG')

h1('6 วัดซ้ำและวัดต่อหลังหยุด')
h2('วัดชิ้นงานเดิมอีกครั้ง')
step(1,'เลือกหมายเลข ALPL ที่ต้องการวัดซ้ำใน Queue')
fig('p094_01','ภาพ 30  เลือกชิ้นงานใน Queue เพื่อวัดซ้ำ',2.65)
step(2,'กด วัดชิ้นนี้อีกรอบ')
fig('p096_01','ภาพ 31  คำสั่งวัดชิ้นงานเดิมอีกครั้ง',2.65)
step(3,'เลือกโหมด Trigger ตรวจชิ้นงาน แล้วกด Start เพื่อเริ่มรอบวัดซ้ำ')
fig('p098_01','ภาพ 32  หน้าต่างยืนยันการวัดซ้ำ',2.65)
step(4,'เมื่อวัดเสร็จ ให้ออกจากโหมดวัดซ้ำเพื่อกลับไปยังคิวหลัก')
fig('p100_01','ภาพ 33  ผลหลังวัดซ้ำ',2.65)
h2('วัดต่อหลัง Session หยุด')
step(1,'เลือก วัดเฉพาะชิ้นที่เหลือ หรือ วัดทั้งหมดอีกครั้ง ตามงานที่ต้องการ')
fig('p103_01','ภาพ 34  ตัวเลือกเมื่อ Session หยุดก่อนวัดครบ',2.65)
step(2,'กด Start เพื่อวัดต่อ หรือกด จบการทำงาน หากไม่ต้องการวัดต่อ')
fig('p105_01','ภาพ 35  ปุ่มวัดต่อและจบการทำงาน',2.65)

h1('7 ดูผลย้อนหลัง')
fig('p015_02','ภาพ 5  ตาราง Measurement History ในหน้า Measure',2.55)
step(1,'เลื่อนลงไปที่ Measurement History แล้วเลือกผลวัดที่ต้องการ')
fig('p110_01','ภาพ 36  ตาราง Measurement History',2.7)
step(2,'ตรวจค่าการวัด ผล OK/NG ภาพ และข้อมูลชิ้นงานในหน้ารายละเอียด')
fig('p112_01','ภาพ 37  รายงานผลวัดรายชิ้น',2.8)

h1('8 เพิ่ม แก้ไข และลบข้อมูล')
h2('ALPL Profile')
para('ใช้ Add Part เพื่อเพิ่มข้อมูล ALPL, Edit เพื่อแก้ไข และปุ่มถังขยะเพื่อลบ หาก ALPL มีผลการวัดอยู่ ระบบอาจไม่อนุญาตให้ลบโปรไฟล์นั้น')
fig('p023_01','ภาพ 6  หน้า Edit ส่วน ALPL Profile',2.55)
fig('p116_01','ภาพ 38  ตาราง ALPL Profile',2.7)
fig('p116_02','ภาพ 39  แบบฟอร์มแก้ไขข้อมูล ALPL',2.7)
h2('Measurement History ในหน้า Edit')
para('เลือก Edit เพื่อแก้ไขผลวัดที่บันทึกไว้ หรือเลือกปุ่มถังขยะเพื่อลบรายการ หน้านี้ใช้แก้ไขและลบผลวัดที่มีอยู่ ไม่ใช้เพิ่มผลวัดใหม่')
fig('p023_02','ภาพ 7  หน้า Edit ส่วน Measurement History',2.55)
fig('p122_01','ภาพ 40  ตารางผลวัดในหน้า Edit',2.7)
fig('p122_02','ภาพ 41  แบบฟอร์มแก้ไขผลวัด',2.7)
h2('Lookup Table และ Trash')
para('Lookup Table ใช้จัดการค่าที่แสดงในรายการเลือกของแบบฟอร์ม ส่วน Trash ใช้ตรวจรายการที่ลบและกู้คืนเมื่อจำเป็น')
fig('p024_01','ภาพ 8  หน้า Edit ส่วนตารางอ้างอิง',2.55)
fig('p024_02','ภาพ 9  หน้า Edit ส่วน Trash',2.55)
fig('p124_01','ภาพ 42  ตัวเลือกตารางใน Lookup Table',1.8)
note('ข้อควรระวัง','ตรวจ ALPL และข้อมูลผลวัดให้ตรงรายการก่อนกด Save หรือลบข้อมูล')

h1('9 ส่งออกข้อมูล')
fig('p030_01','ภาพ 10  หน้า Export',2.55)
step(1,'เปิดเมนู Export แล้วเลือกรูปแบบ CSV, PDF หรือ Excel')
step(2,'เลือก Template ที่กำหนดคอลัมน์หรือรูปแบบรายงาน')
step(3,'กำหนดตัวกรองข้อมูลที่ต้องการส่งออก และตรวจตัวอย่างก่อนดาวน์โหลด')
step(4,'กด Download แล้วตรวจไฟล์ที่บันทึกในเครื่อง')
note('หากไม่พบข้อมูล','ตรวจประเภท Template ตัวกรอง และช่วงวันที่ก่อนลองส่งออกอีกครั้ง')

h1('10 ปัญหาที่พบบ่อย')
h2('Raspberry Pi ไม่ Online')
para('ตรวจว่า Server เป็น Online ก่อน เพราะเมื่อเว็บขาดการเชื่อมต่อ สถานะ Raspberry Pi บนจออาจยังเป็นค่าก่อนหน้า จากนั้นตรวจสายไฟและเครือข่ายของ Raspberry Pi หากยังไม่กลับมา Online ให้แจ้งผู้ดูแลระบบ')
h2('กด Start ไม่ได้')
para('ตรวจว่า Session เป็น STOPPED, Server และ Raspberry Pi เป็น Online แล้วตรวจว่าบันทึก Part Entry ครบทุกช่องที่บังคับและเลือกวิธี Trigger แล้ว')
h2('เกิดปัญหาระหว่างวัด')
para('หยุดการวัด ตรวจสถานะการเชื่อมต่อ และบันทึกหมายเลข ALPL ล่าสุดที่วัดสำเร็จ เมื่อระบบกลับมาพร้อม ให้เลือกวัดเฉพาะชิ้นที่เหลือหรือเริ่มวัดใหม่ตามความเหมาะสม')
h2('ไม่มีผลวัดหรือภาพล่าสุด')
para('ตรวจว่า Session ยังเป็น RUNNING และใช้วิธี Trigger ตรงกับที่ตั้งไว้ หากเป็น Manual ให้กด Trigger แล้วรอผลตอบรับ หากยังไม่ปรากฏผล ให้แจ้งผู้ดูแลระบบ')
h2('ส่งออกแล้วไม่มีข้อมูล')
para('ตรวจตัวกรองข้อมูล ช่วงวันที่ และ Template แล้วทดลองแสดงตัวอย่างก่อนดาวน์โหลดอีกครั้ง')

h1('11 คำศัพท์')
for term,meaning in [
  ('ALPL','หมายเลขระบุชิ้นงาน Alignment Plate'),
  ('Session','รอบการวัดที่เริ่มด้วย Start และสิ้นสุดเมื่อหยุดหรือวัดครบ'),
  ('Tolerance','ค่าความคลาดเคลื่อนที่ยอมรับได้ตามเกณฑ์ของชิ้นงาน'),
  ('OK','ผลวัดผ่านเกณฑ์ที่ตั้งไว้'),
  ('NG','ผลวัดไม่ผ่านเกณฑ์อย่างน้อยหนึ่งรายการ'),
  ('Queue','ลำดับชิ้นงานที่จะวัดใน Session'),
  ('Trigger','คำสั่งให้เครื่องเริ่มวัดชิ้นงาน'),
]:
  q=para(); q.add_run(term+'  ').bold=True; q.add_run(meaning)

h1('12 ผู้ดูแลระบบ')
para('นายภูเบศ จิตภิลัย  แผนก MNT  Teams: Phubet Jitpilai')
para('นายภูบดินทร์ มีหอม  แผนก MNT  Teams: Phubadine Mehom')

d.save(OUT)
print(OUT)
