"""Import a power_bi_gateway.py workbook into an empty MySQL database.

Run on the destination PC:
    python import_powerbi_excel.py "C:\\Data\\tmx_export.xlsx"
    python import_powerbi_excel.py "C:\\Data\\tmx_export.xlsx" --dry-run

Dependencies: pip install pymysql openpyxl python-dotenv
DB_HOST/DB_PORT/DB_NAME/DB_USER/DB_PASSWORD come from the project .env,
environment, or command-line connection options. A missing password is prompted.
No tables are created, cleared, or overwritten. Foreign-key checks stay enabled.
"""
from __future__ import annotations

import argparse
import ast
import getpass
import json
import math
import os
import sys
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path


ALIASES = {"parts_specifications": "parts"}
INTEGER_TYPES = {"tinyint", "smallint", "mediumint", "int", "bigint"}


@dataclass
class Column:
    name: str
    kind: str
    extra: str = ""


@dataclass
class Plan:
    sheet: str
    table: str
    columns: list[Column]
    count: int


def quote(name: str) -> str:
    return "`" + name.replace("`", "``") + "`"


def read_schema(conn, database):
    """Use destination metadata, never hard-code current table/column lists."""
    with conn.cursor() as cur:
        cur.execute("SELECT TABLE_NAME, ENGINE FROM information_schema.TABLES "
                    "WHERE TABLE_SCHEMA=%s AND TABLE_TYPE='BASE TABLE'", (database,))
        engines = dict(cur.fetchall())
        cur.execute("SELECT TABLE_NAME, COLUMN_NAME, DATA_TYPE, EXTRA "
                    "FROM information_schema.COLUMNS WHERE TABLE_SCHEMA=%s "
                    "ORDER BY TABLE_NAME, ORDINAL_POSITION", (database,))
        columns = {table: [] for table in engines}
        for table, name, kind, extra in cur.fetchall():
            if table in columns:
                columns[table].append(Column(name, kind.lower(), extra or ""))
        cur.execute("SELECT TABLE_NAME, REFERENCED_TABLE_NAME, REFERENCED_TABLE_SCHEMA "
                    "FROM information_schema.KEY_COLUMN_USAGE "
                    "WHERE TABLE_SCHEMA=%s AND REFERENCED_TABLE_NAME IS NOT NULL", (database,))
        dependencies = {table: set() for table in engines}
        for child, parent, schema in cur.fetchall():
            if schema != database:
                raise ValueError(f"{child} references another database: {schema}.{parent}")
            dependencies[child].add(parent)
    return engines, columns, dependencies


def table_order(tables, dependencies):
    remaining = set(tables)
    ordered = []
    while remaining:
        ready = sorted(t for t in remaining if not (dependencies.get(t, set()) & remaining))
        if not ready:
            raise ValueError("Circular/self-referencing foreign keys need a dedicated migration: "
                             + ", ".join(sorted(remaining)))
        ordered.extend(ready)
        remaining.difference_update(ready)
    return ordered


def value_for_db(value, column):
    if value is None:
        return None
    if isinstance(value, str) and value.endswith(" …[ตัดแล้ว]"):
        raise ValueError("Cell was truncated by the exporter; use a full DB dump instead")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Non-finite numeric value")
    if column.kind == "json":
        if not isinstance(value, str):
            value = json.dumps(value, ensure_ascii=False, allow_nan=False)
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError:
            # Older exporter serializes Python dictionaries with str(dict).
            try:
                parsed = ast.literal_eval(value)
            except (ValueError, SyntaxError) as exc:
                raise ValueError("Invalid or truncated JSON") from exc
        return json.dumps(parsed, ensure_ascii=False, allow_nan=False)
    if column.kind in INTEGER_TYPES:
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, (int, float)) and abs(value) >= 10**15:
            raise ValueError("Excel numeric cell exceeds 15 digits; original ID may have lost precision")
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value():
            raise ValueError(f"Expected an integer, got {value!r}")
        return int(number)
    if column.kind in {"decimal", "numeric"}:
        return Decimal(str(value))
    if column.kind in {"binary", "varbinary", "tinyblob", "blob", "mediumblob", "longblob", "bit"}:
        raise ValueError("Binary columns cannot be restored reliably from this Excel export")
    return value


def data_rows(sheet, columns):
    """Keep Excel blanks as SQL NULL, strings such as 'NA' as literal strings."""
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
        if all(cell.value is None for cell in cells):
            continue
        result = []
        for cell, column in zip(cells, columns):
            try:
                if cell.data_type in {"f", "e"}:
                    raise ValueError("Formula or Excel error is not a raw exported value")
                result.append(value_for_db(cell.value, column))
            except Exception as exc:
                raise ValueError(f"Sheet {sheet.title}, row {row_number}, "
                                 f"column {column.name}: {exc}") from exc
        yield row_number, tuple(result)


def make_plan(workbook, conn, database):
    engines, columns, dependencies = read_schema(conn, database)
    sheets_to_tables = {}
    for table in columns:
        sheet = ALIASES.get(table, table)[:31]
        if sheet in sheets_to_tables:
            raise ValueError(f"Ambiguous Excel sheet name: {sheet}")
        sheets_to_tables[sheet] = table
    plans = {}
    for sheet in workbook:
        if sheet.title == "_meta":
            continue
        table = sheets_to_tables.get(sheet.title)
        if table is None:
            raise ValueError(f"No matching destination table for sheet {sheet.title}")
        if (engines[table] or "").lower() != "innodb":
            raise ValueError(f"{table} must use InnoDB for all-or-nothing import")
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if not header or any(not isinstance(name, str) for name in header):
            raise ValueError(f"Missing/invalid column header in sheet {sheet.title}")
        if len(set(header)) != len(header):
            raise ValueError(f"Duplicate columns in sheet {sheet.title}")
        actual = {column.name: column for column in columns[table]}
        if set(header) != set(actual):
            raise ValueError(f"Schema mismatch for {table}: "
                             f"missing={sorted(set(actual)-set(header))}, "
                             f"unknown={sorted(set(header)-set(actual))}")
        selected = [actual[name] for name in header]
        if any("GENERATED" in c.extra.upper() for c in selected):
            raise ValueError(f"Generated columns in {table} require a dedicated import")
        # Preflight the entire file before writing any table.
        count = sum(1 for _ in data_rows(sheet, selected))
        plans[table] = Plan(sheet.title, table, selected, count)
    if not plans:
        raise ValueError("Workbook has no data table sheets")
    with conn.cursor() as cur:
        for table in plans:
            cur.execute(f"SELECT 1 FROM {quote(table)} LIMIT 1")
            if cur.fetchone() is not None:
                raise ValueError(f"Destination table {table} is not empty; no data was changed")
    return [plans[table] for table in table_order(plans, dependencies)]


def import_rows(workbook, conn, plans, batch_size=250):
    """One transaction; keep primary IDs and enforce every foreign key."""
    conn.begin()
    try:
        with conn.cursor() as cur:
            # Fail on invalid/out-of-range data rather than silently clipping it.
            cur.execute("SELECT @@SESSION.sql_mode")
            modes = set(filter(None, cur.fetchone()[0].split(",")))
            modes.add("STRICT_ALL_TABLES")
            cur.execute("SET SESSION sql_mode=%s", (",".join(sorted(modes)),))
            for plan in plans:
                cur.execute(f"SELECT 1 FROM {quote(plan.table)} LIMIT 1")
                if cur.fetchone() is not None:
                    raise ValueError(f"Destination {plan.table} is no longer empty")
                names = ", ".join(quote(c.name) for c in plan.columns)
                placeholders = ", ".join(["%s"] * len(plan.columns))
                sql = f"INSERT INTO {quote(plan.table)} ({names}) VALUES ({placeholders})"
                batch = []
                first_row = None
                for row_number, values in data_rows(workbook[plan.sheet], plan.columns):
                    if not batch:
                        first_row = row_number
                    batch.append(values)
                    if len(batch) >= batch_size:
                        insert_batch(cur, sql, batch, plan, first_row, row_number)
                        batch = []
                if batch:
                    insert_batch(cur, sql, batch, plan, first_row, row_number)
                cur.execute(f"SELECT COUNT(*) FROM {quote(plan.table)}")
                if cur.fetchone()[0] != plan.count:
                    raise ValueError(f"Row count mismatch after inserting {plan.table}")
                print(f"Inserted {plan.table}: {plan.count:,} rows (not committed yet)")
        conn.commit()
    except BaseException:
        conn.rollback()
        raise


def insert_batch(cur, sql, batch, plan, first, last):
    try:
        cur.executemany(sql, batch)
    except Exception as exc:
        raise RuntimeError(f"Insert failed: sheet {plan.sheet}, table {plan.table}, "
                           f"Excel rows {first}-{last}: {exc}") from exc


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("excel", nargs="?", help="Path to tmx_export.xlsx; prompted if omitted")
    parser.add_argument("--env-file", type=Path)
    parser.add_argument("--host")
    parser.add_argument("--port", type=int)
    parser.add_argument("--database")
    parser.add_argument("--user")
    parser.add_argument("--batch-size", type=int, default=250)
    parser.add_argument("--dry-run", action="store_true", help="Check schema/values/empty tables; do not insert")
    args = parser.parse_args(argv)
    if args.batch_size < 1:
        parser.error("--batch-size must be positive")
    from dotenv import load_dotenv
    import pymysql
    from openpyxl import load_workbook

    load_dotenv(args.env_file or Path(__file__).resolve().parents[1] / ".env")
    excel = Path((args.excel or input("Excel file path: ")).strip().strip('"')).expanduser().resolve()
    if not excel.is_file() or excel.suffix.lower() != ".xlsx":
        raise ValueError(f"Expected an existing .xlsx file: {excel}")
    database = args.database or os.getenv("DB_NAME", "tmx_db")
    host = args.host or os.getenv("DB_HOST", "127.0.0.1")
    port = args.port or int(os.getenv("DB_PORT", "3306"))
    user = args.user or os.getenv("DB_USER", "pc_user")
    password = os.getenv("DB_PASSWORD")
    if password is None:
        password = getpass.getpass(f"MySQL password for {user}@{host}: ")
    print(f"Destination: {host}:{port}/{database} as {user}")
    print(f"Workbook: {excel}")
    workbook = load_workbook(excel, read_only=True, data_only=False)
    conn = None
    try:
        conn = pymysql.connect(host=host, port=port, database=database, user=user,
                               password=password, charset="utf8mb4", autocommit=False,
                               connect_timeout=10)
        plans = make_plan(workbook, conn, database)
        print("Import order:")
        for plan in plans:
            print(f"  {plan.sheet} -> {plan.table}: {plan.count:,} rows")
        if args.dry_run:
            conn.rollback()
            print("Dry run passed. No data was inserted; FK/unique constraints are checked during actual import.")
        else:
            import_rows(workbook, conn, plans, args.batch_size)
            print(f"COMMITTED: {sum(p.count for p in plans):,} rows across {len(plans)} tables")
    finally:
        workbook.close()
        if conn:
            conn.close()
    return 0


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"IMPORT FAILED: {exc}\nNo import transaction was committed.", file=sys.stderr)
        raise SystemExit(1)
